import pandas as pd
import openpyxl
from services.timing_extractor import extract_timings_from_text, normalize_time_range
from services.validator import validate_timetable_structure
import re

def extract_from_excel(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active
    
    timing_row_idx = None
    days_col_idx = None
    
    days_of_week = ['monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday']
    days_abbrev = [d[:3] for d in days_of_week]
    
    matrix = []
    for r in sheet.iter_rows(values_only=True):
        matrix.append(list(r))
        
    periods = []
    period_cols = {} 
    
    for r_idx, row in enumerate(matrix):
        time_matches = 0
        temp_periods = []
        for c_idx, cell in enumerate(row):
            if cell and isinstance(cell, str):
                extracted = extract_timings_from_text(cell)
                if extracted:
                    time_matches += 1
                    norm = normalize_time_range(extracted[0])
                    if norm:
                        temp_periods.append({
                            'col_idx': c_idx,
                            'id': f"P{time_matches}",
                            'start': norm['start'],
                            'end': norm['end']
                        })
        if time_matches >= 3:
            timing_row_idx = r_idx
            for tp in temp_periods:
                periods.append({
                    'id': tp['id'],
                    'start': tp['start'],
                    'end': tp['end']
                })
                period_cols[tp['col_idx']] = tp['id']
            break
            
    for r_idx, row in enumerate(matrix):
        for c_idx, cell in enumerate(row):
            if cell and isinstance(cell, str):
                c_lower = cell.strip().lower()
                if c_lower in days_of_week or c_lower in days_abbrev:
                    days_col_idx = c_idx
                    break
        if days_col_idx is not None:
            break
            
    if timing_row_idx is None or days_col_idx is None:
        raise ValueError("Could not reliably detect the timetable timing row or days column.")
        
    merged_map = {}
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        val = sheet.cell(row=min_row, column=min_col).value
        for r in range(min_row - 1, max_row):
            for c in range(min_col - 1, max_col):
                if (r, c) != (min_row - 1, min_col - 1):
                    merged_map[(r, c)] = {
                        'value': val,
                        'origin': (min_row - 1, min_col - 1),
                        'span_cols': max_col - min_col + 1
                    }

    extracted_days = []
    timetable_data = {}
    
    for r_idx in range(timing_row_idx + 1, len(matrix)):
        row = matrix[r_idx]
        day_cell = row[days_col_idx]
        if not day_cell:
            continue
            
        day_str = str(day_cell).strip().capitalize()
        if day_str.lower() in days_of_week or day_str.lower()[:3] in [d[:3] for d in days_of_week]:
            day_full = next((d.capitalize() for d in days_of_week if d.startswith(day_str.lower()[:3])), day_str)
            extracted_days.append(day_full)
            timetable_data[day_full] = {}
            
            for c_idx, p_id in period_cols.items():
                cell_val = row[c_idx]
                is_merged = False
                span = 1
                
                if (r_idx, c_idx) in merged_map:
                    cell_val = merged_map[(r_idx, c_idx)]['value']
                    is_merged = True
                    span = merged_map[(r_idx, c_idx)]['span_cols']
                else:
                    for merged_range in sheet.merged_cells.ranges:
                        if merged_range.bounds[1] - 1 == r_idx and merged_range.bounds[0] - 1 == c_idx:
                            is_merged = True
                            span = merged_range.bounds[2] - merged_range.bounds[0] + 1
                            break
                            
                if cell_val:
                    cell_str = str(cell_val).strip().replace('\n', ' ')
                    cell_type = 'lecture'
                    
                    cell_str_clean = cell_str.replace(' ', '')
                    if 'lab' in cell_str.lower():
                        cell_type = 'lab'
                    elif 'break' in cell_str_clean.lower() or 'lunch' in cell_str_clean.lower():
                        cell_type = 'break'
                        
                    if is_merged and span > 1:
                        affected_pids = []
                        for offset in range(span):
                            if (c_idx + offset) in period_cols:
                                affected_pids.append(period_cols[c_idx + offset])
                        
                        if affected_pids and affected_pids[0] == p_id:
                            timetable_data[day_full][p_id] = {
                                'subject': cell_str,
                                'type': cell_type,
                                'periods': affected_pids
                            }
                        elif not affected_pids:
                            timetable_data[day_full][p_id] = {
                                'subject': cell_str,
                                'type': cell_type
                            }
                    else:
                        timetable_data[day_full][p_id] = {
                            'subject': cell_str,
                            'type': cell_type
                        }
                        
    return {
        'periods': periods,
        'days': extracted_days,
        'timetable': timetable_data
    }
